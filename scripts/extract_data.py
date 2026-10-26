"""
Exporta a aba "Recebimentos" do Controle_Recebimentos.xlsx para um JSON
"achatado" — foi esse o formato usado como dado de entrada (const DATA)
do painel interativo (index.html).

Clientes sem Região/Categoria preenchida (ou que não estão no Cadastro de
Clientes) entram no total, mas caem no grupo "Não identificado" nos dois
campos, igual é feito nos gráficos do dashboard em Excel.

Uso:
    pip install openpyxl
    python scripts/extract_data.py Controle_Recebimentos.xlsx transacoes.json
"""

import json
import sys

from openpyxl import load_workbook

MESES_EXTENSO = [
    "Janeiro", "Fevereiro", "Março", "Abril", "Maio", "Junho",
    "Julho", "Agosto", "Setembro", "Outubro", "Novembro", "Dezembro",
]


def extract(xlsx_path: str, json_path: str) -> None:
    wb = load_workbook(xlsx_path, data_only=True)
    rec = wb["Recebimentos"]

    rows = []
    for row in rec.iter_rows(min_row=5, max_col=6):
        data, cliente, valor, forma, regiao, categoria = (c.value for c in row)
        if data is None or cliente is None:
            continue
        rows.append({
            "date": data.strftime("%Y-%m-%d"),
            "month": MESES_EXTENSO[data.month - 1],
            "cliente": cliente,
            "valor": round(float(valor), 2),
            "forma": forma,
            "regiao": regiao or "Não identificado",
            "categoria": categoria or "Não identificado",
        })

    with open(json_path, "w", encoding="utf-8") as f:
        json.dump(rows, f, ensure_ascii=False, indent=2)

    total = sum(r["valor"] for r in rows)
    print(f"{len(rows)} lançamentos exportados para {json_path} (total R$ {total:,.2f})")


if __name__ == "__main__":
    if len(sys.argv) != 3:
        print("uso: python scripts/extract_data.py Controle_Recebimentos.xlsx transacoes.json")
        raise SystemExit(1)
    extract(sys.argv[1], sys.argv[2])
